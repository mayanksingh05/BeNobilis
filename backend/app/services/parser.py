import io
import logging
from typing import BinaryIO
import pdfplumber
import docx

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_stream: BinaryIO) -> str:
    """
    Extract selectable text from a PDF file with layout-awareness
    and robust error handling.
    """
    text_chunks = []

    try:
        with pdfplumber.open(file_stream) as pdf:
            if not pdf.pages:
                raise ValueError("The uploaded PDF has no pages.")

            for page_idx, page in enumerate(pdf.pages):
                page_text = page.extract_text(layout=True)

                # Fallback to standard text extraction if layout extraction is empty
                if not page_text or not page_text.strip():
                    page_text = page.extract_text()

                if page_text and page_text.strip():
                    text_chunks.append(page_text.strip())

    except ValueError:
        raise
    except Exception as e:
        err_msg = str(e).lower()
        if "password" in err_msg or "encrypted" in err_msg:
            raise ValueError("The PDF is password-protected. Please upload an unprotected copy.")
        logger.error(f"Failed to parse PDF: {e}")
        raise ValueError(f"Unable to read PDF document: {e}")

    full_text = "\n\n".join(text_chunks).strip()

    if len(full_text) < 30:
        raise ValueError(
            "Unable to extract readable text. The PDF may be a scanned image or screenshot. "
            "Please upload a document with selectable text."
        )

    return full_text


def extract_text_from_docx(file_stream: BinaryIO) -> str:
    """
    Extract text from a DOCX file including both paragraphs and tables.
    """
    text_chunks = []

    try:
        doc = docx.Document(file_stream)

        # 1. Extract paragraphs
        for para in doc.paragraphs:
            content = para.text.strip()
            if content:
                text_chunks.append(content)

        # 2. Extract tables (many resumes format skills/education in tables)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))

    except Exception as e:
        logger.error(f"Failed to parse DOCX: {e}")
        raise ValueError(f"Unable to read DOCX document: {e}")

    full_text = "\n".join(text_chunks).strip()

    if len(full_text) < 30:
        raise ValueError(
            "The uploaded DOCX document is empty or contains insufficient text."
        )

    return full_text


def extract_text(file_stream: BinaryIO, filename: str = "", content_type: str = "") -> str:
    """
    Unified extraction entrypoint for PDF and DOCX files.
    """
    fname = (filename or "").lower()
    ctype = (content_type or "").lower()

    if fname.endswith(".docx") or "wordprocessingml" in ctype:
        return extract_text_from_docx(file_stream)
    elif fname.endswith(".pdf") or "pdf" in ctype:
        return extract_text_from_pdf(file_stream)
    else:
        # Attempt PDF first, then fallback to DOCX
        try:
            return extract_text_from_pdf(file_stream)
        except Exception:
            file_stream.seek(0)
            return extract_text_from_docx(file_stream)